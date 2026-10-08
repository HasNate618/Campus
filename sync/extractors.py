"""Pure file→text extraction, shared by the sync engine and the agent's
`extract_file` tool.

No DB, no Config object, and no side effects beyond the Office→PDF cache file
that `sync.convert` already writes beside the source. Callers decide what to do
with the text (write a `.md` sibling, return it to the model, discard it).

Page markers use the same `<!-- page N -->` format as `agent.citations.PAGE_RE`,
so page-addressed reads keep working on anything extracted here.
"""
from __future__ import annotations

from pathlib import Path

#: Suffixes `extract_to_text` knows how to turn into text.
EXTRACTABLE_SUFFIXES = frozenset({
    ".pdf", ".docx", ".pptx", ".doc",
    ".md", ".markdown", ".txt", ".rst",
    ".csv", ".tsv", ".json", ".yaml", ".yml",
    ".html", ".htm", ".xml", ".ics",
    ".py", ".js", ".ts", ".tsx", ".java", ".c", ".cpp", ".h", ".go",
    ".rs", ".sh", ".sql", ".css",
})

_OFFICE_SUFFIXES = frozenset({".docx", ".pptx"})


def _looks_like_data_table(t) -> bool:
    """Truth-table-style grids (rows of short cells) — NOT the slide-deck
    layout boxes PyMuPDF's table finder also reports (those have huge cells:
    header/footer text columns). Short-cell filter keeps only real data tables."""
    try:
        if t.row_count < 2 or t.col_count < 2:
            return False
        ext = t.extract()
        if not ext:
            return False
        cells = [str(c).strip() for row in ext for c in row if c is not None]
        if not cells:
            return False
        return max(len(c) for c in cells) <= 30
    except Exception:
        return False


def _pdf_text(path: Path) -> str | None:
    """Digital PDFs carry an embedded text layer — PyMuPDF reads it without
    OCR. Scans come back empty (None); the caller may route them to an
    external parser or leave them unextracted. Data tables are appended as
    markdown, matching the sync engine's original extraction."""
    import pymupdf

    parts: list[str] = []
    doc = pymupdf.open(path)
    try:
        page_no = 0
        for page in doc:
            page_no += 1
            text = page.get_text()
            if not text.strip():
                continue
            parts.append(f"<!-- page {page_no} -->\n{text.rstrip()}")
            try:
                found = page.find_tables()
                tables = [t for t in (found.tables if found else [])
                          if _looks_like_data_table(t)]
            except Exception:
                tables = []
            for t in tables:
                parts.append("")
                parts.append(t.to_markdown())
    finally:
        doc.close()
    return "\n".join(parts) or None


def _doc_text(path: Path) -> str | None:
    """Legacy `.doc` via antiword; `.docx` via python-docx. (`sync` routes
    `.docx` through LibreOffice→PDF for better structure; this direct path is
    the fallback the tool uses when soffice is unavailable.)"""
    if path.suffix.lower() == ".docx":
        try:
            import docx  # python-docx
        except ImportError:
            return None
        text = "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
        return text or None
    import subprocess

    try:
        out = subprocess.run(
            ["antiword", str(path)], capture_output=True, text=True, timeout=60
        )
    except (FileNotFoundError, subprocess.TimeoutExpired):
        return None
    return out.stdout if out.returncode == 0 and out.stdout.strip() else None


def extract_to_text(path: Path, *, max_bytes: int = 20 * 1024 * 1024) -> str | None:
    """Return the text of `path`, or None when it can't be extracted.

    None means "no text" (unsupported suffix, missing file, oversize, or an
    empty/scanned document) — never an exception. Callers surface that to the
    user/model rather than treating it as an error.
    """
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in EXTRACTABLE_SUFFIXES or not path.is_file():
        return None
    try:
        if path.stat().st_size > max_bytes:
            return None
        if suffix == ".pdf":
            return _pdf_text(path)
        if suffix in _OFFICE_SUFFIXES:
            # LibreOffice → PDF → text. convert_office_to_pdf returns None when
            # soffice is absent or the file can't be converted; for .docx fall
            # back to python-docx so the tool still works without soffice.
            from sync.convert import convert_office_to_pdf

            pdf = convert_office_to_pdf(path)
            if pdf is not None:
                return _pdf_text(pdf)
            return _doc_text(path) if suffix == ".docx" else None
        if suffix == ".doc":
            return _doc_text(path)
        return path.read_text(encoding="utf-8", errors="replace")
    except Exception:
        return None
